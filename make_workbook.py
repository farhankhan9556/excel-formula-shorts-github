#!/usr/bin/env python3
from __future__ import annotations
from datetime import datetime
from pathlib import Path
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from formula_library import FORMULAS

START_ROW = 5

def translate_formula(formula: str) -> str:
    if not formula or not formula.startswith("="):
        return formula
    return Translator(formula, origin="A2").translate_formula(f"A{START_ROW}")

def convert_value(header: str, value):
    if value == "" or value is None:
        return value
    if isinstance(value, (int, float, datetime)):
        return value
    s = str(value)
    h = header.lower()
    if "date" in h:
        for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(s, fmt)
            except ValueError:
                pass
    if any(x in h for x in ("id", "name", "dept", "department", "status", "action")):
        return s
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+\.\d+", s):
        return float(s)
    return s

def result_number_format(result: str) -> str:
    s = str(result)
    if "%" in s:
        return "0%"
    if re.fullmatch(r"\d{2}-[A-Za-z]{3}-\d{4}", s):
        return "dd-mmm-yyyy"
    if re.fullmatch(r"-?\d+\.\d+", s):
        return "0.00"
    return "General"

def create_workbook(item: dict, output: Path):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Excel Tip"
    ws.sheet_view.showGridLines = True
    ws.freeze_panes = "A5"

    ws["A1"] = "Learn Verse — Excel Quick Tip"
    ws["A1"].font = Font(size=18, bold=True, color="217346")
    ws["A2"] = item["problem"]
    ws["A2"].font = Font(size=11, italic=True, color="555555")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(3, len(item["headers"])))

    for c, header in enumerate(item["headers"], 1):
        cell = ws.cell(4, c, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="217346")
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(item["rows"], START_ROW):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, convert_value(item["headers"][c - 1], value))

    formula_col = len(item["headers"]) + 2
    result_col = formula_col + 1
    demo_col = result_col + 1

    for c, title in ((formula_col, "Formula"), (result_col, "Outcome"), (demo_col, "Live Entry")):
        cell = ws.cell(4, c, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    formula = translate_formula(item["formula"])
    ws.cell(START_ROW, formula_col, formula)
    ws.cell(START_ROW, result_col, item["result"])
    ws.cell(START_ROW, demo_col, "")
    ws.cell(START_ROW + 1, demo_col, "Type here")
    ws.cell(START_ROW + 1, demo_col).font = Font(italic=True, color="888888")

    ws.cell(START_ROW, formula_col).number_format = result_number_format(item["result"])
    if "date" in " ".join(item["headers"]).lower():
        for c, h in enumerate(item["headers"], 1):
            if "date" in h.lower():
                for r in range(START_ROW, START_ROW + len(item["rows"])):
                    ws.cell(r, c).number_format = "dd-mmm-yyyy"

    thin = Side(style="thin", color="D9E1F2")
    max_row = START_ROW + max(0, len(item["rows"]) - 1)
    for row in ws.iter_rows(min_row=4, max_row=max_row + 1, min_col=1, max_col=demo_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center")

    for c in range(1, demo_col + 1):
        values = [str(ws.cell(r, c).value or "") for r in range(1, max_row + 2)]
        width = max(14, min(34, max(len(v) for v in values) + 3))
        ws.column_dimensions[get_column_letter(c)].width = width

    ws.auto_filter.ref = f"A4:{get_column_letter(len(item['headers']))}{max_row}"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(output)

    check = load_workbook(output, read_only=True, data_only=False)
    actual = check["Excel Tip"].cell(START_ROW, formula_col).value
    check.close()
    if actual != formula:
        raise RuntimeError(f"Formula verification failed: {actual!r} != {formula!r}")

    print(f"Workbook created: {output}")
    print(f"Formula: {formula}")
    print(f"Expected outcome: {item['result']}")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python make_workbook.py <index> <output.xlsx>")
    create_workbook(FORMULAS[int(sys.argv[1])], Path(sys.argv[2]))
